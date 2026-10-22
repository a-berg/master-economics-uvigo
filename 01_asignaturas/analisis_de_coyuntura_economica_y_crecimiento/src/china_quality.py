"""Independent semantic checks, source comparisons and quality report for China."""
import hashlib,json,math,re
from pathlib import Path
import polars as pl
from china_silver import ROOT,OUT,manifest
from pipeline import write_json

def main():
    d=pl.read_parquet(ROOT/'datos/silver/china.parquet');cat=pl.read_csv(OUT/'catalogo_china.csv',infer_schema_length=None)
    assert not d.select(pl.struct('series_id','year').is_duplicated().any()).item()
    assert d['value'].is_finite().all() and d['value'].null_count()==0
    assert d.filter(pl.col('analytic') & (pl.col('year')>2025)).is_empty()
    assert d.filter(pl.col('analytic') & pl.col('unit').str.contains('pending|not_explicit')).is_empty()
    prov=json.loads((OUT/'provenance.json').read_text());hashes={m['sha256'] for m in prov}
    for m in prov:assert hashlib.sha256((ROOT/m['artifact']).read_bytes()).hexdigest()==m['sha256'],m['artifact']
    values={(r['series_id'],r['year']):r['value'] for r in d.iter_rows(named=True)}
    def val(s,y):return values.get((s,y))
    def close(a,b,abs_tol=1e-8):return math.isclose(a,b,rel_tol=1e-10,abs_tol=abs_tol)
    for r in d.iter_rows(named=True):
        if r['source']=='derived':
            assert r['input_series'] and json.loads(r['dimensions'])['input_hashes']
            assert set(json.loads(r['dimensions'])['input_hashes'])<=hashes
        else:assert r['bronze_sha256'] in hashes
    wdi_checks=0
    for m in prov:
        if not m['id'].startswith('china_wdi_data_'):continue
        obj=json.loads((ROOT/m['artifact']).read_text())
        if len(obj)!=2:continue
        for r in obj[1]:
            if r['value'] is None:continue
            assert val('WDI.'+r['indicator']['id'],int(r['date']))==float(r['value']);wdi_checks+=1
    # PWT original-cell checks on the full selected coverage (not just sentinels).
    from openpyxl import load_workbook
    pwt_checks=0
    for source,prefix in [('pwt110','PWT.'),('pwt110_na','NA.'),('pwt110_capital','CAP.')]:
        m=manifest(source);b=load_workbook(ROOT/m['artifact'],read_only=True,data_only=True);s=b['Data'];s.reset_dimensions();it=s.iter_rows(values_only=True);header=next(it)
        for rr in it:
            r=dict(zip(header,rr))
            if r.get('countrycode')!='CHN':continue
            for key,v in r.items():
                if (prefix+key,r['year']) in values:
                    assert values[prefix+key,r['year']]==float(v);pwt_checks+=1
        b.close()
    ages=[];labor=[];formulas=0
    for y in sorted(set(d['year'])):
        pop=val('WDI.SP.POP.TOTL',y);young=val('WDI.SP.POP.0014.TO',y);work=val('WDI.SP.POP.1564.TO',y);old=val('WDI.SP.POP.65UP.TO',y)
        if all(v is not None for v in [pop,young,work,old]):
            diff=pop-young-work-old;ages.append(dict(year=y,difference_persons=diff));assert abs(diff)<=3
        lf=val('ILO.EAP_2EAP_SEX_AGE_NB.XA:2070',y);emp=val('ILO.EMP_2EMP_SEX_AGE_NB.XA:2070',y);unemp=val('ILO.UNE_2UNE_SEX_AGE_NB.XA:2070',y)
        if all(v is not None for v in [lf,emp,unemp]):
            residual=lf-emp-unemp;labor.append(dict(year=y,residual_thousand=residual));assert abs(residual)<=.0021
        # Arithmetic identities recomputed independently.
        for sid,expected in [
            ('DER.pop_15plus',pop-young if pop is not None and young is not None else None),
            ('DER.hours_total',val('PWT.emp',y)*val('PWT.avh',y) if val('PWT.emp',y) is not None and val('PWT.avh',y) is not None else None),
            ('DER.gdp_deflator',100*val('WDI.NY.GDP.MKTP.CN',y)/val('WDI.NY.GDP.MKTP.KN',y) if val('WDI.NY.GDP.MKTP.CN',y) is not None else None)]:
            if expected is not None:assert close(val(sid,y),expected);formulas+=1
    pl.DataFrame(ages).write_csv(OUT/'age_identity.csv');pl.DataFrame(labor).write_csv(OUT/'ilo_identity.csv')
    # UN table 2.6 only. Counter AND explicit no-data marker required.
    m=manifest('china_un26_presence');text=(ROOT/m['artifact']).read_text()
    counter=re.search(r'id="spanRecordCountB">([\d,]+)</span>',text)
    assert counter,'UN response ambiguous; no absence conclusion allowed'
    un_count=int(counter.group(1).replace(',',''))
    if un_count==0:assert 'No data was found matching your query' in text
    write_json(OUT/'un26_presence.json',dict(table='2.6',group=206,country_code=156,classification='ISIC Rev.4',records=un_count,evidence=m,limit='Esta consulta y variante; no prueba ausencia universal. No se usa AMA.'))
    # Contrasts are evidence of consistency/revisions, NEVER independent validation.
    comparisons=[]
    for y in range(1952,2026):
        for label,a,b,scale,meaning in [
            ('NBS_GDP_vs_WDI','NBS.GDP.CURRENT','WDI.NY.GDP.MKTP.CN',1e8,'Misma magnitud nominal tras escala; ediciones y revisiones pueden diferir; dependencia nacional no descartada'),
            ('NBS_GDP_vs_PWT_NA_scale_hypothesis','NBS.GDP.CURRENT','NA.v_gdp',100.,'Hipótesis escala NA=millones CNY, solo diagnóstico; no se convierte la silver ni acredita equivalencia'),
            ('NBS_M2_vs_WDI_broad_money','NBS.M2.YEAR_END','WDI.FM.LBL.BMNY.CN',1e8,'M2 frente a broad money; verificar definición/fecha; WDI redistribuye FMI y datos nacionales, no independiente')]:
            aa=val(a,y);bb=val(b,y)
            if aa is None or bb is None:continue
            comparisons.append(dict(comparison=label,year=y,left_series=a,right_series=b,left_original=aa,right_original=bb,left_multiplier=scale,left_scaled=aa*scale,
                difference=aa*scale-bb,relative_difference_pct=100*(aa*scale/bb-1) if bb else None,
                independence='NO_ACREDITADA',interpretation=meaning))
    cmp=pl.DataFrame(comparisons);cmp.write_csv(OUT/'source_comparisons.csv')
    cmp.group_by('comparison').agg(pl.len().alias('n'),pl.col('relative_difference_pct').abs().max().alias('max_abs_relative_pct'),pl.col('relative_difference_pct').median().alias('median_relative_pct')).sort('comparison').write_csv(OUT/'source_comparison_summary.csv')
    # Same-source annual growth with consecutive years and positive endpoints.
    growth=d.sort('series_id','year').with_columns(pl.col('year').shift().over('series_id').alias('previous_year'),pl.col('value').shift().over('series_id').alias('previous_value'))
    growth=growth.filter((pl.col('year')==pl.col('previous_year')+1)&(pl.col('value')>0)&(pl.col('previous_value')>0)).with_columns((100*(pl.col('value')/pl.col('previous_value')-1)).alias('growth_pct'))
    growth.select('series_id','year','growth_pct','quality','analytic').write_csv(OUT/'annual_growth.csv')
    growth.filter(pl.col('growth_pct').abs()>25).select('series_id','year','growth_pct','quality').write_csv(OUT/'large_changes_for_review.csv')
    # Required 1960–2023 growth accounting coverage and endpoint decomposition.
    periods=[(1960,1989),(1989,2000),(2000,2008),(2008,2019),(2019,2023),(1960,2023)]
    accounting=[]
    for start,end in periods:
        inputs=['PWT.rgdpna','PWT.rnna','PWT.emp']
        assert all(val(s,y) is not None for s in inputs for y in range(start,end+1))
        changes=[100*math.log(val(s,end)/val(s,start))/(end-start) for s in inputs]
        g,k,l=changes;residual=g-k/3-2*l/3
        accounting.append(dict(start=start,end=end,gdp_log_growth=g,capital_contribution=k/3,employment_contribution=2*l/3,solow_residual=residual,alpha=1/3,n_years=end-start+1,note='Mismos periodos temas 6–7; residuo no identifica tecnología; PWT estimaciones y PIB oficial CHN'))
    pl.DataFrame(accounting).write_csv(OUT/'growth_accounting_1960_2023.csv')
    contracts=pl.read_csv(OUT/'contratos.csv',infer_schema_length=None)
    quality=dict(status='OK_WITH_DOCUMENTED_GAPS',observations=d.height,series=cat.height,analytic_observations=d.filter(pl.col('analytic')).height,
        analytic_series=d.filter(pl.col('analytic'))['series_id'].n_unique(),quarantined_series=cat.filter(~pl.col('analytic')).height,
        original_manifest_hash_checks=len(prov),independent_raw_wdi_cell_checks=wdi_checks,independent_raw_pwt_cell_checks=pwt_checks,
        age_identity_checks=len(ages),ilo_identity_checks=len(labor),formula_checks=formulas,un26_records=un_count,
        contracts=contracts.height,contract_statuses=contracts.group_by('status').len().sort('status').to_dicts(),
        caveat='Controles de integridad y coherencia; no certifican precisión de la medición ni independencia ni comparabilidad plena.')
    write_json(OUT/'quality.json',quality);print(json.dumps(quality,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
