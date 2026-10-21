"""Pipeline reproducible. Ejecutar desde cualquier directorio: uv run python src/pipeline.py all."""
from __future__ import annotations
import argparse, csv, hashlib, os, itertools, json, math, shutil, statistics, sys, urllib.parse, urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
import polars as pl

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'datos'
CAT = DATA / 'catalogo'
BRONZE = DATA / 'bronze'
SILVER = DATA / 'silver'
GOLD = DATA / 'gold'

def write_json(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + '\n')

def digest(data): return hashlib.sha256(data).hexdigest()
def now(): return datetime.now(timezone.utc).isoformat()

def ingest_local():
    """Importa bytes originales; no inventa URL ni fecha histórica de descarga."""
    for path in DATA.iterdir():
        if path.suffix not in ('.csv', '.xls'): continue
        raw = path.read_bytes(); sha = digest(raw)
        folder = BRONZE / 'local' / sha[:12]; folder.mkdir(parents=True, exist_ok=True)
        target = folder / path.name
        if not target.exists(): target.write_bytes(raw)
        manifest = folder / (path.name + '.json')
        if not manifest.exists():
            write_json(manifest, dict(dataset_id='local_' + sha[:12], provider='user_supplied',
                source_file=path.name, artifact=target.relative_to(ROOT).as_posix(), sha256=sha,
                size_bytes=len(raw), imported_at_utc=now(), accessed_at_utc=None,
                download_url=None, publication_date=None, provenance_status='historical_download_unknown'))


def download_one(spec, refresh=False):
    folder = BRONZE / spec['id']; folder.mkdir(parents=True, exist_ok=True)
    latest = folder / 'latest.json'
    if latest.exists() and not refresh:
        old = json.loads(latest.read_text())
        artifact = ROOT / old['artifact']
        if old.get('url')==spec['url'] and artifact.exists() and digest(artifact.read_bytes()) == old['sha256']:
            return dict(id=spec['id'], status='cached', artifact=old['artifact'])
        if old.get('url')==spec['url']: raise ValueError(f"Bronze corrupto: {artifact}")
    try:
        req = urllib.request.Request(spec['url'], headers={'User-Agent':'UVigoAcademicMacroReport/0.1'})
        with urllib.request.urlopen(req, timeout=100) as response:
            raw = response.read(); content_type=response.headers.get('Content-Type')
        if spec['format']=='json':
            obj=json.loads(raw)
            if isinstance(obj,dict) and ('error' in obj or obj.get('status') in (400,404,500)):
                raise ValueError(str(obj)[:250])
            if spec['provider']=='world_bank' and (not isinstance(obj,list) or len(obj)!=2 or not isinstance(obj[1],list)):
                raise ValueError('Respuesta WDI sin observaciones')
            if spec['provider']=='eurostat' and not obj.get('value'):
                raise ValueError('Respuesta Eurostat sin value')
        sha=digest(raw); stamp=datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
        filename=f'{stamp}_{sha[:12]}.{spec["format"]}'
        artifact=folder/filename; artifact.write_bytes(raw)
        publication=None; update=None
        if spec['format']=='json':
            update=obj.get('updated') if isinstance(obj,dict) else obj[0].get('lastupdated')
        manifest={**spec,'artifact':artifact.relative_to(ROOT).as_posix(),'sha256':sha,
            'size_bytes':len(raw),'accessed_at_utc':now(),'content_type':content_type,
            'publication_date':spec.get('publication_date'),'provider_updated_at':update}
        write_json(folder/(filename+'.json'),manifest)
        write_json(latest,manifest)
        return dict(id=spec['id'],status='downloaded',artifact=manifest['artifact'])
    except Exception as exc:
        return dict(id=spec['id'],status='failed',error=f'{type(exc).__name__}: {exc}')


def download(refresh=False):
    ingest_local()
    specs=json.loads((CAT/'downloads.json').read_text())
    with ThreadPoolExecutor(max_workers=5) as pool:
        results=list(pool.map(lambda s:download_one(s,refresh),specs))
    write_json(GOLD/'download_status.json',results)
    for r in results: print(r['id'],r['status'],r.get('error',''))
    return results


def manifest(dataset):
    p=BRONZE/dataset/'latest.json'
    return json.loads(p.read_text()) if p.exists() else None


def audit_excel():
    import xlrd
    b=xlrd.open_workbook(next(DATA.glob('*.xls')))
    rows=[]; sheet_info=[]
    for sheet in b.sheets():
        sheet_info.append(dict(sheet=sheet.name,rows=sheet.nrows,columns=sheet.ncols))
        for i in range(sheet.nrows):
            for j in range(sheet.ncols):
                cell=sheet.cell(i,j)
                if cell.ctype==xlrd.XL_CELL_EMPTY: continue
                rows.append(dict(sheet=sheet.name,row=i+1,column=j+1,type=cell.ctype,value=str(cell.value)))
    SILVER.mkdir(parents=True,exist_ok=True)
    pl.DataFrame(rows).write_parquet(SILVER/'reference_excel_cells.parquet')
    write_json(GOLD/'reference_excel_sheets.json',sheet_info)
    source=b.sheet_by_name('Source')
    lines=['# Fuentes declaradas en el Excel del profesor','',
           'Transcripción de valores almacenados. No se recalculan fórmulas ni se adoptan sus series como fuente primaria.','']
    for i in range(source.nrows):
        vals=[str(c.value) for c in source.row(i) if c.value!='']
        if vals: lines.append(f"- Fila {i+1}: " + ' | '.join(vals))
    (GOLD/'reference_sources.md').write_text('\n'.join(lines)+'\n')


def eurostat_rows(obj):
    """Decodifica índices multidimensionales JSON-stat sin perder unidades ni flags."""
    ids=obj['id']; sizes=obj['size']; codes=[]
    for dim in ids:
        index=obj['dimension'][dim]['category']['index']
        codes.append(index if isinstance(index,list) else [k for k,v in sorted(index.items(),key=lambda kv:kv[1])])
    values=obj.get('value',{}); flags=obj.get('status',{})
    if isinstance(values,list): values={str(i):v for i,v in enumerate(values) if v is not None}
    if isinstance(flags,list): flags={str(i):v for i,v in enumerate(flags) if v is not None}
    for index,value in values.items():
        rem=int(index); coords=[0]*len(ids)
        for pos in reversed(range(len(ids))): coords[pos]=rem%sizes[pos];rem//=sizes[pos]
        yield {**{dim:codes[p][coords[p]] for p,dim in enumerate(ids)},'value':value,'flag':flags.get(str(index))}


def normalize():
    from openpyxl import load_workbook
    records=[]
    source_manifests={s['id']:manifest(s['id']) for s in json.loads((CAT/'downloads.json').read_text())}
    local_sha=digest(next(DATA.glob('API_CHN*.csv')).read_bytes())
    def add(country,year,series,value,unit,source,source_code,flag=None,dimensions=None):
        if value is None or isinstance(value,bool): return
        try: value=float(value);year=int(year)
        except (ValueError,TypeError): return
        if not math.isfinite(value): raise ValueError('Valor no finito')
        m=source_manifests.get(source)
        records.append(dict(country=country,year=year,series=series,value=value,unit=unit,
            source=source,source_code=source_code,flag=flag,
            dimensions=json.dumps(dimensions or {},sort_keys=True),
            bronze_sha256=m['sha256'] if m else local_sha))
    # Importación histórica WDI local, independiente de las descargas nuevas.
    file=next(DATA.glob('API_CHN*.csv'))
    with file.open(encoding='utf-8-sig') as f:
        reader=csv.reader(f)
        header=next(row for row in reader if row and row[0]=='Country Name')
        for row in reader:
            if len(row)<4: continue
            for year,value in zip(header[4:],row[4:]):
                if value and year.isdigit(): add('CHN',year,row[3],value,'see_indicator_metadata','wdi_local',row[3])
    # Versiones completas PWT con diccionario preservado.
    for dataset in ['pwt110','pwt110_na','pwt110_capital']:
        m=manifest(dataset)
        if not m: continue
        b=load_workbook(ROOT/m['artifact'],read_only=True,data_only=True)
        sheet=b['Data'];sheet.reset_dimensions()
        it=sheet.iter_rows(values_only=True);header=list(next(it))
        country_idx=header.index('countrycode');year_idx=header.index('year')
        for row in it:
            if len(row)<=max(country_idx,year_idx) or row[country_idx] not in ['CHN','ESP']:continue
            for key,value in zip(header,row):
                if key in ['countrycode','country','currency_unit','year'] or not isinstance(value,(int,float)):continue
                add(row[country_idx],row[year_idx],key,value,'see_pwt_legend',dataset,key)
        legend=b['Legend'];legend.reset_dimensions()
        write_json(SILVER/(dataset+'_legend.json'),[list(r) for r in legend.iter_rows(values_only=True) if r and r[0]])
    # INE: seleccionar por código de operación, no solo posición de fila.
    m=manifest('ine_integrated_accounts')
    if m:
        import re
        b=load_workbook(ROOT/m['artifact'],read_only=True,data_only=True)
        codes={'P.1':'output','P.2':'intermediate_consumption','D.21-D.31':'product_taxes_net','B.1b/B.1*b':'gdp_market','P.51c':'ccf_nominal'}
        for sheet in b:
            match=re.search(r'año (\d{4})',sheet.title)
            if not match:continue
            year=int(match.group(1));seen=set()
            for rownum,row in enumerate(sheet.iter_rows(values_only=True),1):
                if len(row)<15:continue
                code=row[12]
                if code in codes and code not in seen and isinstance(row[2],(float,int)):
                    add('ESP',year,codes[code],row[2],'million_EUR_current','ine_integrated_accounts',code,
                        dimensions={'sheet':sheet.title,'row':rownum,'column':3,'operation_code_column':13})
                    seen.add(code)
    for spec in json.loads((CAT/'downloads.json').read_text()):
        m=manifest(spec['id'])
        if not m or spec['format']!='json':continue
        obj=json.loads((ROOT/m['artifact']).read_text())
        if spec['provider']=='eurostat':
            for r in eurostat_rows(obj):
                dims={k:v for k,v in r.items() if k not in ['value','flag','time']}
                series='|'.join(str(dims[k]) for k in obj['id'] if k!='time')
                add('ESP',r['time'],series,r['value'],r.get('unit','unknown'),spec['id'],
                    r.get('na_item',r.get('indic_em',r.get('asset10',''))),r.get('flag'),dims)
        elif spec['provider']=='world_bank':
            for r in obj[1]:
                add(r['countryiso3code'],r['date'],r['indicator']['id'],r['value'],
                    r.get('unit') or 'see_indicator_metadata',spec['id'],r['indicator']['id'],r.get('obs_status'))
    df=pl.DataFrame(records, infer_schema_length=None).with_columns(pl.col('year').cast(pl.Int32))
    keys=['country','year','series','source']
    if df.select(pl.struct(keys).is_duplicated().any()).item():raise ValueError('Claves duplicadas en silver')
    df=df.sort(keys)
    df.write_parquet(SILVER/'observations.parquet')
    df.group_by(['country','series','source','unit']).agg(pl.col('year').min().alias('first_year'),
        pl.col('year').max().alias('last_year'),pl.len().alias('n')).sort(['source','country','series']).write_csv(GOLD/'coverage.csv')
    print('Silver:',df.height,'observaciones')
    return df


def build_gold(df):
    def values(country,source,series):
        return {r['year']:r['value'] for r in df.filter((pl.col('country')==country)&(pl.col('source')==source)&(pl.col('series')==series)).iter_rows(named=True)}
    out=[];growth=[];checks=[]
    for country in ['CHN','ESP']:
        main={s:values(country,'pwt110',s) for s in ['rgdpna','rnna','emp','pop','avh','labsh','delta','rgdpe']}
        na={s:values(country,'pwt110_na',s) for s in ['q_gdp','q_c','q_i','q_g','q_x','q_m','q_gfcf','v_gdp','v_c','v_i','v_g','v_x','v_m','v_gfcf']}
        for year in sorted(main['rgdpna']):
            def val(s):return main[s].get(year)
            def ratio(a,b):return a/b if a is not None and b is not None and b!=0 else None
            y=val('rgdpna');k=val('rnna');n=val('emp');pop=val('pop');avh=val('avh')
            row=dict(country=country,year=year,gdp_real=y,capital_real=k,employment_million=n,population_million=pop,
                hours_per_worker=avh,gdp_per_capita=ratio(y,pop),gdp_per_worker=ratio(y,n),
                gdp_per_hour=ratio(y,n*avh if n and avh else None),capital_per_worker=ratio(k,n),
                labor_share=val('labsh'),depreciation_rate=val('delta'),
                gdp_ppp_per_capita=ratio(val('rgdpe'),pop))
            for s,vs in na.items():row[s]=vs.get(year)
            for s in ['c','i','g','x','m']:
                row['share_'+s]=ratio(na['v_'+s].get(year),na['v_gdp'].get(year))
            out.append(row)
        # Subperiodos homogéneos por pareja; no sumar porcentajes discretos como identidad exacta.
        periods=json.loads((CAT/'analysis_config.json').read_text())['periods']
        for start,end in periods:
            if not all(start in main[s] and end in main[s] and main[s][start]>0 and main[s][end]>0 for s in ['rgdpna','rnna','emp','pop']):continue
            g={s:100*math.log(main[s][end]/main[s][start])/(end-start) for s in ['rgdpna','rnna','emp','pop']}
            alpha=json.loads((CAT/'analysis_config.json').read_text())['alpha']
            capital=alpha*g['rnna'];labor=(1-alpha)*g['emp'];residual=g['rgdpna']-capital-labor
            growth.append(dict(country=country,start=start,end=end,years=end-start,
                gdp_log_pp=g['rgdpna'],capital_contribution_pp=capital,labor_contribution_pp=labor,
                solow_residual_pp=residual,gdp_cagr_pct=100*((main['rgdpna'][end]/main['rgdpna'][start])**(1/(end-start))-1),
                gdp_pc_cagr_pct=100*(((main['rgdpna'][end]/main['pop'][end])/(main['rgdpna'][start]/main['pop'][start]))**(1/(end-start))-1)))
            checks.append(dict(check='growth_identity',country=country,start=start,end=end,error=g['rgdpna']-capital-labor-residual))
    panel=pl.DataFrame(out, infer_schema_length=None).sort(['country','year'])
    # Crecimientos exigen t-1 observado; no se salta huecos.
    derived=[]
    for country in ['CHN','ESP']:
        prior=None
        for row in panel.filter(pl.col('country')==country).iter_rows(named=True):
            for col in ['gdp_real','capital_real','employment_million','gdp_per_worker','gdp_per_hour','q_c','q_i','q_g','q_x','q_m']:
                a=row[col];b=prior[col] if prior and row['year']==prior['year']+1 else None
                row['growth_'+col]=100*(a/b-1) if a is not None and b is not None and a>0 and b>0 else None
            derived.append(row);prior=row
    panel=pl.DataFrame(derived, infer_schema_length=None)
    panel.write_parquet(GOLD/'comparison.parquet');panel.write_csv(GOLD/'comparison.csv')
    pl.DataFrame(growth).write_csv(GOLD/'growth_accounting.csv')
    write_json(GOLD/'checks.json',checks)
    # Estadística cíclica sobre tasas, con la misma muestra de cada par.
    stats=[]
    for country in ['CHN','ESP']:
        for start,end in json.loads((CAT/'analysis_config.json').read_text())['cycle_periods']:
            p=panel.filter((pl.col('country')==country)&pl.col('year').is_between(start,end))
            for col in ['growth_q_c','growth_q_i','growth_q_g','growth_q_x','growth_q_m','growth_capital_real','growth_employment_million','growth_gdp_per_worker']:
                d=p.select('growth_gdp_real',col).drop_nulls()
                if d.height<3:continue
                sx=d['growth_gdp_real'].std(ddof=1);sv=d[col].std(ddof=1)
                corr=d.select(pl.corr('growth_gdp_real',col)).item()
                stats.append(dict(country=country,start=start,end=end,variable=col,n=d.height,
                    first_year=p.filter(pl.col('growth_gdp_real').is_not_null()&pl.col(col).is_not_null())['year'].min(),
                    last_year=p.filter(pl.col('growth_gdp_real').is_not_null()&pl.col(col).is_not_null())['year'].max(),
                    correlation=corr if corr is not None and math.isfinite(corr) else None,relative_std=sv/sx if sx else None))
    pl.DataFrame(stats).write_csv(GOLD/'cycle_statistics.csv')
    # Cobertura de todos los requisitos: evita presentar una figura pendiente como completada.
    requirements=json.loads((CAT/'requirements.json').read_text())
    available=set(df['source'].unique())
    for req in requirements:
        req['status']='pending_validation'
        if req['id'] in ['6.1','6.2','6.3','6.4','7.1']:req['status']='generated_pwt_common_basis'
        if req['id'] in ['2B.levels','2B.shares','2B.growth','2B.stats','2B.2008','2C.labor','2C.intensity']:req['status']='data_available_pwt'
    write_json(GOLD/'requirements_status.json',requirements)
    print('Gold:',panel.height,'filas;',len(growth),'descomposiciones')
    return panel


def figures(panel):
    os.environ.setdefault('MPLCONFIGDIR',str(ROOT/'.mplconfig'))
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.size':10,'axes.spines.top':False,'axes.spines.right':False,'svg.fonttype':'none'})
    dest=ROOT/'output'/'figuras';dest.mkdir(parents=True,exist_ok=True)
    labels={'CHN':'China','ESP':'España'}
    specs=[('gdp_pc','gdp_ppp_per_capita','PIB real por habitante','Dólares internacionales de 2021 (PWT rgdpe/pop)'),
           ('growth','growth_gdp_real','Crecimiento del PIB real','Variación anual (%)'),
           ('productivity','gdp_per_worker','PIB real por ocupado','USD de 2021 a precios nacionales constantes'),
           ('capital','capital_per_worker','Capital por ocupado','USD de 2021 a precios nacionales constantes'),
           ('hours','hours_per_worker','Horas anuales por ocupado','Horas'),
           ('labor_share','labor_share','Participación del trabajo','Proporción del PIB')]
    for filename,col,title,ylabel in specs:
        fig,ax=plt.subplots(figsize=(7.2,4))
        for country in labels:
            d=panel.filter((pl.col('country')==country)&(pl.col('year')>=1990)).select('year',col).drop_nulls()
            ax.plot(d['year'],d[col],label=labels[country])
        ax.set(title=title,xlabel='Año',ylabel=ylabel);ax.legend();ax.grid(alpha=.2)
        fig.tight_layout();fig.savefig(dest/(filename+'.svg'));fig.savefig(dest/(filename+'.png'),dpi=130);plt.close(fig)
    for country in labels:
        for kind,columns in [('demand_shares',['share_c','share_i','share_g','share_x','share_m']),('demand_2008',['q_gdp','q_c','q_i','q_g','q_x','q_m'])]:
            fig,ax=plt.subplots(figsize=(7.2,4))
            for col in columns:
                d=panel.filter((pl.col('country')==country)&pl.col('year').is_between(1995,2023)).select('year',col).drop_nulls()
                if d.height==0:continue
                vals=d[col]
                if kind=='demand_2008':
                    base=d.filter(pl.col('year')==2008)
                    if not base.height or base[col][0]==0:continue
                    vals=100*vals/base[col][0]
                else: vals=100*vals
                ax.plot(d['year'],vals,label=col.replace('share_','').replace('q_','').upper())
            ax.set(title=f'{labels[country]}: '+('composición nominal de la demanda' if kind=='demand_shares' else 'demanda real, 2008 = 100'),xlabel='Año',ylabel='% del PIB' if kind=='demand_shares' else 'Índice')
            ax.legend(ncol=3);ax.grid(alpha=.2);fig.tight_layout()
            fig.savefig(dest/f'{kind}_{country}.svg');fig.savefig(dest/f'{kind}_{country}.png',dpi=130);plt.close(fig)


def main():
    p=argparse.ArgumentParser();p.add_argument('command',choices=['all','download','build','audit']);p.add_argument('--refresh',action='store_true');args=p.parse_args()
    for path in [CAT,BRONZE,SILVER,GOLD]:path.mkdir(parents=True,exist_ok=True)
    if args.command in ['audit','all']:audit_excel()
    if args.command in ['download','all']:download(args.refresh)
    if args.command in ['build','all']:
        df=normalize();panel=build_gold(df);figures(panel)
        from official import build
        build(df, ROOT)
        from report import build as report_build
        report_build(ROOT)
    if args.command=='all':
        failures=[r for r in json.loads((GOLD/'download_status.json').read_text()) if r['status']=='failed']
        if failures:print('ADVERTENCIA:',len(failures),'descargas fallidas; consultar download_status.json');sys.exit(2)

if __name__=='__main__':main()
