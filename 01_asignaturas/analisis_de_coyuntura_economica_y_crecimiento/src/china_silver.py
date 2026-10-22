"""Build China in isolation from archived originals; no network and no splicing."""
from __future__ import annotations
import hashlib,json,math,re
from collections import defaultdict
from pathlib import Path
import polars as pl
from openpyxl import load_workbook
from pipeline import write_json
from china_acquire import INDICATORS
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/revision_china'
SILVER=ROOT/'datos/silver'

def manifest(id):
    p=ROOT/'datos/bronze'/id/'latest.json'
    if not p.exists():return None
    m=json.loads(p.read_text());a=ROOT/m['artifact']
    if hashlib.sha256(a.read_bytes()).hexdigest()!=m['sha256']:raise ValueError(f'Hash mismatch: {a}')
    return m

def wdi_unit(code):
    if code.endswith('.CN'):return 'CNY_current'
    if code.endswith('.KN'):return 'CNY_constant_provider_base'
    if code.startswith('SP.POP') or code=='SL.TLF.TOTL.IN':return 'persons'
    if code=='FP.CPI.TOTL':return 'index_2010_100'
    if code=='SI.POV.GINI':return 'index_0_100'
    if code=='PA.NUS.FCRF':return 'CNY_per_USD_annual_average'
    if code.startswith('SP.DYN'):return 'per_1000_population'
    return 'percent'

# Definition, selected candidate, Spanish references (conceptual, never automatic equality).
CONCEPTS={
'gdp_nominal':('PIB a precios de mercado corrientes','WDI.NY.GDP.MKTP.CN','UVGD'),
'gdp_real':('PIB en volumen; base nacional','WDI.NY.GDP.MKTP.KN','OVGD'),
'c_nominal':('Consumo final hogares e ISFLSH corriente','WDI.NE.CON.PRVT.CN','UCPH'),
'c_real':('Consumo hogares; volumen','NA.q_c','OCPH'),
'g_nominal':('Consumo final administraciones corriente','WDI.NE.CON.GOVT.CN','UCTG'),
'g_real':('Consumo administraciones; volumen','NA.q_g','OCTG'),
'i_nominal':('Formación bruta de capital corriente','WDI.NE.GDI.TOTL.CN','UITT'),
'i_real':('Inversión PWT; alcance FBC por revisar','NA.q_i','OITT'),
'gfcf_real':('Formación bruta de capital fijo volumen','NA.q_gfcf','OIGT'),
'inventories_real':('Variación de existencias volumen; sin resta de cadenas','','OIST'),
'ccf_real':('Consumo de capital fijo en volumen; no tasa delta','','OKCT'),
'capital_real':('Stock neto PWT a precios nacionales constantes','PWT.rnna','OKND'),
'x_nominal':('Exportaciones bienes y servicios corrientes','WDI.NE.EXP.GNFS.CN','UXGS'),
'x_real':('Exportaciones bienes y servicios volumen','NA.q_x','OXGS'),
'm_nominal':('Importaciones bienes y servicios corrientes','WDI.NE.IMP.GNFS.CN','UMGS'),
'm_real':('Importaciones bienes y servicios volumen','NA.q_m','OMGS'),
'output':('Producción P.1; economía total; corriente','NBS.IO.GROSS_OUTPUT','INE.output'),
'intermediate_consumption':('Consumo intermedio P.2; economía total; corriente','NBS.IO.INTERMEDIATE_INPUTS','INE.intermediate_consumption'),
'gdp_factor_cost':('PIB menos D2 más D3; coste factores','','UYGD'),
'compensation':('Remuneración asalariados D1; excluye autónomos','NBS.COMPENSATION.DOMESTIC.USES','UWCD'),
'gross_operating_surplus':('Excedente bruto; renta mixta identificada separadamente','','UOGD'),
'taxes':('D2 impuestos producción e importación brutos','','UTVT'),
'subsidies':('D3 subvenciones producción e importación','','UYVT'),
'input_output':('Matriz flujos intersectoriales, tres sectores; año estructural','','tabla_2D.2'),
'cpi':('IPC nacional; no IPCA','WDI.FP.CPI.TOTL','ZCPIN'),
'gdp_deflator':('Cociente PIB corriente/volumen mismo proveedor','DER.gdp_deflator','PVGD'),
'employment':('Personas ocupadas PWT; no equivalentes','PWT.emp','NETN;PWT.emp'),
'hours':('Horas totales anuales PWT empleo por horas medias','DER.hours_total','PWT.avh;NETN'),
'population':('Población total WDI','WDI.SP.POP.TOTL','NPTN'),
'pop_0_14':('Población 0–14 inclusive','WDI.SP.POP.0014.TO','NPCN'),
'pop_15plus':('Población 15+; complemento de 0–14','DER.pop_15plus','NPAN;NPON'),
'pop_65plus':('Población 65+','WDI.SP.POP.65UP.TO','NPON'),
'labor_force':('Activos 15+ modelizados OIT; no EPA','ILO.EAP_2EAP_SEX_AGE_NB.XA:2070','INE.EPA.ACTIVOS'),
'unemployment':('Desempleados 15+ modelizados; nivel directo','ILO.UNE_2UNE_SEX_AGE_NB.XA:2070','INE.EPA.PARADOS'),
'registered_unemployment':('Paro registrado; ámbito urbano por documentar','NBS.REGISTERED_UNEMPLOYED.URBAN','SEPE.REGISTERED.UNEMPLOYED.ANNUAL_MEAN'),
'survey_unemployment':('Desempleados encuesta nacional; universo ILO conservado','ILO.UNE_TUNE_SEX_AGE_NB.BA:16044','INE.EPA.PARADOS'),
'gini':('Gini encuestas WDI; años observados sin interpolación','WDI.SI.POV.GINI','WDI.SI.POV.GINI'),
'equivalent_employment':('Aproximación docente horas anuales/1760','DER.equivalent_employment','FETD'),
}

def build():
    OUT.mkdir(parents=True,exist_ok=True)
    rows=[];catalog={};used={};wdi_coverage=[]
    def add(sid,year,value,unit,m,name,quality='PUBLICADA',analytic=True,dimensions=None,flag='',formula='',inputs='',note='',upstream='Cadena nacional específica pendiente; no independencia acreditada'):
        if value is None:return
        value=float(value)
        if not math.isfinite(value):raise ValueError(sid)
        year=int(year)
        if m:used[m['sha256']]=m
        rows.append(dict(country='CHN',year=year,series_id=sid,value=value,unit=unit,
            source=m['id'] if m else 'derived',source_code=sid.split('.',1)[-1],
            bronze_sha256=m['sha256'] if m else '',flag=flag or '',dimensions=json.dumps(dimensions or {},sort_keys=True),
            quality=quality,analytic=analytic and year<=2025,formula=formula,input_series=inputs))
        if sid not in catalog:
            catalog[sid]=dict(series_id=sid,name=name,unit=unit,source=m['id'] if m else 'derived',
                source_url=m.get('url','') if m else '',accessed_at_utc=m.get('accessed_at_utc') if m else None,
                bronze_sha256=m['sha256'] if m else '',quality=quality,analytic=analytic,
                formula=formula,input_series=inputs,method_note=note,upstream=upstream,
                independence='NO_ACREDITADA',comparability='REVISAR_UNIVERSO_Y_BASE; no comparar niveles monetarios locales entre países')
    # Fresh WDI edition, isolated from historical user download and wdi_common.
    for code in INDICATORS:
        m=manifest('china_wdi_data_'+code.lower().replace('.','_'))
        meta=manifest('china_wdi_meta_'+code.lower().replace('.','_'))
        if not m:continue
        obj=json.loads((ROOT/m['artifact']).read_text())
        valid=isinstance(obj,list) and len(obj)==2 and isinstance(obj[1],list)
        if valid and obj[0].get('pages')!=1:raise ValueError('WDI incomplete pagination')
        metadata=json.loads((ROOT/meta['artifact']).read_text()) if meta else []
        md=metadata[1][0] if len(metadata)==2 and metadata[1] else {}
        if meta:used[meta['sha256']]=meta
        count=0
        for r in obj[1] if valid else []:
            if r['value'] is None:continue
            assert r['countryiso3code']=='CHN'
            count+=1
            add('WDI.'+code,r['date'],r['value'],wdi_unit(code),m,md.get('name',r['indicator']['value']),
                dimensions={'indicator':code,'provider_lastupdated':obj[0].get('lastupdated')},flag=r.get('obs_status'),
                note=md.get('sourceNote',''),upstream=md.get('sourceOrganization','No documentada'))
        used[m['sha256']]=m
        wdi_coverage.append(dict(code=code,n=count,status='PUBLICADA' if count else 'SIN_VALORES' if valid else 'CODIGO_RECHAZADO_API',metadata=json.dumps(md,ensure_ascii=False)))
    # Re-read PWT original workbooks, don't rely on silver as a source.
    for dataset,prefix in [('pwt110','PWT.'),('pwt110_na','NA.'),('pwt110_capital','CAP.')]:
        m=manifest(dataset);book=load_workbook(ROOT/m['artifact'],read_only=True,data_only=True)
        sheet=book['Legend'];sheet.reset_dimensions()
        legend={str(r[0]):str(r[1]) for r in sheet.iter_rows(values_only=True) if len(r)>1 and r[0] and r[1]}
        sheet=book['Data'];sheet.reset_dimensions();it=sheet.iter_rows(values_only=True);header=next(it)
        for rr in it:
            r=dict(zip(header,rr))
            if r.get('countrycode')!='CHN':continue
            for k,v in r.items():
                if k in ['year','countrycode','country','currency_unit'] or not isinstance(v,(float,int)):continue
                if prefix=='PWT.' and k not in ['rgdpna','rnna','rgdpe','pop','emp','avh','labsh','delta','rtfpna','hc']:continue
                name=legend.get(k,k)
                if prefix=='PWT.':
                    unit={'rgdpna':'million_USD_2021_national_constant','rnna':'million_USD_2021_national_constant','rgdpe':'million_USD_2021_chained_PPP','pop':'million_persons','emp':'million_persons','avh':'hours_per_engaged_person_year','labsh':'share','delta':'rate','rtfpna':'index_2021_1','hc':'index' }[k]
                    analytic=True;quality='ESTIMACION_PWT'
                else:
                    unit='native_currency_current_scale_pending' if prefix=='CAP.' or k.startswith('v_') else 'native_currency_constant_scale_pending'
                    if k in ['emp','pop']:unit='million_persons'
                    elif k=='avh':unit='hours_per_engaged_person_year'
                    elif k in ['xr','xr2']:unit='national_currency_per_USD'
                    elif prefix=='CAP.' and k.startswith('Ksh_'):unit='share'
                    elif prefix=='CAP.' and k.startswith(('Ip_','Np_','Kp_')):unit='price_index_reference_pending'
                    analytic='pending' not in unit
                    quality='CANDIDATA_BASE_INDICE_PENDIENTE' if 'price_index' in unit else 'CANDIDATA_ESCALA_PENDIENTE' if not analytic else 'ESTIMACION_PWT'
                add(prefix+k,r['year'],v,unit,m,name,quality,analytic,
                    note='Leyenda preservada. NA/CAP: niveles monetarios en cuarentena hasta acreditar escala y valoración; tasas e índices pueden explorarse sin certificar equivalencia. PWT CHN usa PIB oficial; CH2/Wu se mantiene separado.',
                    upstream='PWT 11; PIB CHN oficial según What is new; empleo requiere i_emp. Sin validación independiente NBS.')
        book.close()
    # ILO: preserve source, age, notes, status. Never stitch official estimates/LFS/census.
    m=manifest('china_ilo_chn_annual')
    if m:
        d=pl.read_csv(ROOT/m['artifact'],infer_schema_length=10000)
        inds=pl.read_csv(ROOT/manifest('china_ilo_dic_indicator')['artifact'])
        names=dict(zip(inds['indicator'],inds['indicator.label']))
        select=d.filter((pl.col('ref_area')=='CHN') & (pl.col('sex')=='SEX_T') & (pl.col('classif1')=='AGE_YTHADULT_YGE15') & pl.col('classif2').is_null() & pl.col('indicator').is_in([
            'EAP_2EAP_SEX_AGE_NB','EAP_2WAP_SEX_AGE_RT','EMP_2EMP_SEX_AGE_NB','EMP_2WAP_SEX_AGE_RT','UNE_2UNE_SEX_AGE_NB','UNE_2EAP_SEX_AGE_RT',
            'EMP_TEMP_SEX_AGE_NB','UNE_TUNE_SEX_AGE_NB','UNE_DEAP_SEX_AGE_RT']))
        for r in select.iter_rows(named=True):
            sid='ILO.'+r['indicator']+'.'+r['source'];model=r['source']=='XA:2070'
            add(sid,r['time'],r['obs_value'],'thousand_persons' if r['indicator'].endswith('_NB') else 'percent',m,names[r['indicator']],
                'MODELIZADA_OIT' if model else 'PUBLICADA_UNIVERSO_NACIONAL',True,
                dimensions={k:r[k] for k in ['source','sex','classif1','classif2','note_source','note_indicator','note_classif']},flag=r['obs_status'],
                note='13th ICLS; edad codificada 15+. Fuente nacional puede tener ajustes: conservar notas. Proyecciones 2026+ excluidas de vista analítica. Encuesta, censo y estimación oficial separados.',
                upstream='OIT modelización' if model else 'Fuente nacional distribuida por OIT; no medición independiente')
        d.group_by('indicator','source','sex','classif1','classif2').agg(pl.col('time').min().alias('first'),pl.col('time').max().alias('last'),pl.len().alias('n')).sort('indicator','source','sex','classif1','classif2').write_csv(OUT/'ilo_inventory.csv')
        for dic in ['indicator','source','classif1','classif2','sex','obs_status','note_source','note_indicator','note_classif']:
            mm=manifest('china_ilo_dic_'+dic)
            if mm:used[mm['sha256']]=mm
    # NBS fallback: transcribed cells only, no automatic interpolation of annual history.
    if manifest('china_nbs_03_21'):
        from china_nbs_extract import extract
        for r in extract():
            m=manifest('china_nbs_'+r['table'])
            add(r['series_id'],r['year'],r['value'],r['unit'],m,r['series_id'],
                'ALTERNATIVA_NBS_COBERTURA_LIMITADA',True,
                dimensions={'table':r['table'],'extraction':r['extraction']},note=r['note'],
                upstream='NBS' if r['table']!='18_01' else 'People Bank of China vía NBS')
        for key,value in [('GROSS_OUTPUT',3381462.2),('INTERMEDIATE_INPUTS',2080497.0),('VALUE_ADDED',1300965.2),('CCF.CURRENT',195362.6)]:
            add('NBS.IO.'+key,2023,value,'native_IO_table_unit_scale_not_explicit_in_image',manifest('china_nbs_03_21'),key,
                'ALTERNATIVA_ESTRUCTURAL_ESCALA_PENDIENTE',False,dimensions={'table':'3-21','valuation':'producer_prices_2023'},
                note='Una observación estructural. VAB a precios de productor; no igualar PIB mercado. Total publicado; el VAB suma componentes con redondeo. Escala no explícita en JPG.',upstream='NBS input-output 2023')
    # Explicit formulas; same-year intersections only.
    def derive(sid,inputs,fn,unit,name,note=''):
        by={i:{r['year']:r for r in rows if r['series_id']==i} for i in inputs}
        years=sorted(set.intersection(*(set(v) for v in by.values())))
        for year in years:
            source_rows=[by[i][year] for i in inputs];values=[r['value'] for r in source_rows]
            if any(v is None for v in values):continue
            value=fn(*values)
            if value is None:continue
            hashes=set()
            for r in source_rows:
                if r['bronze_sha256']:hashes.add(r['bronze_sha256'])
                hashes.update(json.loads(r['dimensions']).get('input_hashes',[]))
            add(sid,year,value,unit,None,name,'DERIVADA',all(r['analytic'] for r in source_rows),
                dimensions={'input_hashes':sorted(hashes)},formula=name,inputs=';'.join(inputs),note=note)
    derive('DER.gdp_deflator',['WDI.NY.GDP.MKTP.CN','WDI.NY.GDP.MKTP.KN'],lambda a,b:100*a/b if b>0 else None,'index_provider_reference','100 * PIB nominal WDI / PIB volumen WDI')
    derive('DER.pop_15plus',['WDI.SP.POP.TOTL','WDI.SP.POP.0014.TO'],lambda a,b:a-b,'persons','Población total - población 0–14')
    derive('DER.hours_total',['PWT.emp','PWT.avh'],lambda a,b:a*b,'million_hours','emp * avh','Estimación coherente PWT; no observación independiente de horas.')
    derive('DER.equivalent_employment',['DER.hours_total'],lambda a:a*1e6/1760,'persons_equivalent','horas totales * 1e6 / 1760','Aproximación docente; no serie observada.')
    derive('DER.gdp_per_worker',['PWT.rgdpna','PWT.emp'],lambda a,b:a/b if b>0 else None,'USD_2021_national_constant_per_person','rgdpna / emp','Apto crecimiento; no niveles comparativos de bienestar.')
    derive('DER.capital_output',['PWT.rnna','PWT.rgdpna'],lambda a,b:a/b if b>0 else None,'ratio','rnna / rgdpna')
    derive('DER.gdp_pc_ppp',['PWT.rgdpe','PWT.pop'],lambda a,b:a/b if b>0 else None,'USD_2021_PPP_per_person','rgdpe / pop')
    # Within-series normalization cancels the unresolved constant monetary scale.
    # This permits growth exploration, not certification of native monetary levels.
    for var in ['q_gdp','q_c','q_i','q_g','q_x','q_m','q_gfcf']:
        series=[r for r in rows if r['series_id']=='NA.'+var]
        base=next((r['value'] for r in series if r['year']==2008),None)
        if base is None or base<=0:continue
        for r in series:
            add('INDEX.NA.'+var,r['year'],100*r['value']/base,'index_2008_100',None,'PWT NA '+var+' / valor 2008 * 100',
                'DERIVADA_SOLO_CRECIMIENTO',True,dimensions={'input_hashes':[r['bronze_sha256']]},
                formula='100 * NA.'+var+' / NA.'+var+'[2008]',inputs='NA.'+var,
                note='Escala constante se cancela. Misma serie sin empalme. No certifica alcance FBC/ISFLSH ni comparabilidad de niveles con España.')
    df=pl.DataFrame(rows,infer_schema_length=None).sort('series_id','year')
    assert not df.select(pl.struct('series_id','year').is_duplicated().any()).item()
    df.write_parquet(SILVER/'china.parquet');df.write_csv(SILVER/'china.csv')
    analytic=df.filter(pl.col('analytic'))
    analytic.pivot(on='series_id',index='year',values='value').sort('year').write_parquet(SILVER/'china_anual.parquet')
    analytic.pivot(on='series_id',index='year',values='value').sort('year').write_csv(SILVER/'china_anual.csv')
    for sid,c in catalog.items():
        ys=set(df.filter(pl.col('series_id')==sid)['year'].to_list());c.update(first=min(ys),last=max(ys),n=len(ys),missing_1960_2023=';'.join(map(str,sorted(set(range(1960,2024))-ys))))
    cat=pl.DataFrame(list(catalog.values()),infer_schema_length=None).sort('series_id');cat.write_csv(OUT/'catalogo_china.csv');cat.write_parquet(SILVER/'china_catalogo.parquet')
    for path in sorted((ROOT/'datos/bronze').glob('china_*/latest.json')):
        m=manifest(path.parent.name)
        used[m['sha256']]=m
    for id in ['pwt110_whatsnew','pwt110_userguide','pwt110_labor_detail']:
        m=manifest(id)
        if m:used[m['sha256']]=m
    write_json(OUT/'provenance.json',list(used.values()))
    pl.DataFrame(wdi_coverage).write_csv(OUT/'wdi_inventory.csv')
    old=pl.read_parquet(SILVER/'observations.parquet').filter(pl.col('country')=='CHN')
    old.group_by('source','series','unit').agg(pl.col('year').min().alias('first'),pl.col('year').max().alias('last'),pl.len().alias('n'),pl.col('bronze_sha256').first()).sort('source','series').write_csv(OUT/'existing_inventory.csv')
    contracts(df,cat)
    return df,cat

def contracts(df,cat):
    req=json.loads((ROOT/'datos/catalogo/requirements.json').read_text());refs=defaultdict(list)
    for r in req:
        for k in r['inputs']:refs[k].append(r['id'])
    refs['equivalent_employment'].append('4.labor; instructions PDF p.6')
    cr=[]
    for key,(definition,sid,spain) in CONCEPTS.items():
        match=cat.filter(pl.col('series_id')==sid);c=match.row(0,named=True) if match.height else {}
        ys=set(df.filter(pl.col('series_id')==sid)['year'].to_list())
        status='NO_RECUPERADA'
        if c:status='DERIVADA' if sid.startswith('DER.') else 'CANDIDATA_EN_CUARENTENA' if not c['analytic'] else 'ALTERNATIVA_DOCUMENTADA' if sid.startswith(('ILO.','PWT.')) else 'CORRESPONDENCIA_CONCEPTUAL'
        if c and sid.startswith('NBS.'):status='ALTERNATIVA_NBS_COBERTURA_LIMITADA' if c['analytic'] else 'ALTERNATIVA_ESTRUCTURAL_ESCALA_PENDIENTE'
        if key=='input_output' and (OUT/'io_2023_3_sectors.csv').exists():status='ALTERNATIVA_ESTRUCTURAL_2023_ESCALA_PENDIENTE'
        monetary=key.endswith(('_real','_nominal')) or key in ['compensation','gross_operating_surplus','taxes','subsidies','gdp_factor_cost','output','intermediate_consumption','input_output']
        unit='CNY a precios nacionales constantes; base y escala explícitas' if key.endswith('_real') else 'CNY corrientes; escala explícita' if monetary else 'horas anuales totales' if key=='hours' else 'índice (base explícita)' if key in ['cpi','gdp_deflator'] else 'Gini 0–100' if key=='gini' else 'personas'
        universe='China continental; economía total'
        if key in ['labor_force','unemployment','survey_unemployment']:universe='China continental; ambos sexos; 15+; separar encuesta de modelización y conservar universo original'
        elif key=='registered_unemployment':universe='Registro administrativo; separar cobertura urbana y nacional; sin equivalencia automática con SEPE'
        elif key=='employment':universe='Todos los ocupados; distinguir doméstico/nacional, asalariados/autónomos y personas/ETC'
        elif key=='compensation':universe='D1 pagado por unidades productoras domésticas a asalariados; no renta laboral ajustada de autónomos'
        elif key=='c_nominal' or key=='c_real':universe='Hogares e ISFLSH: confirmar inclusión en cada proveedor'
        elif key.startswith('pop') or key=='population':universe='Población residente; grupos por edad definidos, media/cierre por documentar'
        elif key=='gini':universe='Encuesta de hogares: renta/consumo y cobertura específicos pendientes de acreditar para comparación'
        cr.append(dict(concept=key,definition=definition,required_by=';'.join(refs.get(key,[])),input_kind='derivado' if sid.startswith('DER.') else 'input_estructural' if key=='input_output' else 'input_base',
            universe=universe,required_unit=unit,
            required_valuation='volumen nacional constante' if key.endswith('_real') else 'precios corrientes; distinguir mercado/básicos/productor/factores' if monetary else 'no monetaria; índice/medida según definición',frequency='estructural un año' if key=='input_output' else 'anual; no mezclar media y cierre',
            required_period='1960–2023 (tema 7 explícito; máximo observado otros temas; no obligación universal de panel completo)',spain_reference=spain,
            selected_series=sid if c else '',candidate_series=sid,growth_alternative='INDEX.'+sid if sid.startswith('NA.q_') else '',structural_artifact='output/revision_china/io_2023_3_sectors.csv' if key=='input_output' else '',status=status,unit=c.get('unit',''),first=min(ys) if ys else None,last=max(ys) if ys else None,n=len(ys),
            missing_1960_2023=';'.join(map(str,sorted(set(range(1960,2024))-ys))),
            comparability='PENDIENTE; no acredita igualdad con España' if not c or not c['analytic'] else 'CONCEPTUAL; armonizar universo/base; niveles monetarios locales no comparables',
            decision='No imputar; no empalmar; revisar notas específicas del catálogo',source_url=c.get('source_url','')))
    pl.DataFrame(cr,infer_schema_length=None).write_csv(OUT/'contratos.csv')
    # Every Spanish catalogue row accounted for; use existing carefully reviewed AMECO rules as candidate search index.
    from catalogue import rules
    ruleset=rules();mapping=[]
    sp=pl.read_csv(ROOT/'output/revision_espana/catalogo_espana.csv',infer_schema_length=None)
    suffix_to_key={s:k for k,(_,_,ss) in CONCEPTS.items() for s in ss.split(';')}
    for r in sp.iter_rows(named=True):
        suffix=r['series_id'].split('.')[-1];key=suffix_to_key.get(suffix) or suffix_to_key.get(r['series_id'])
        co=next((x for x in cr if x['concept']==key),None)
        candidate=ruleset.get(suffix,{}).get('CHN')
        mapping.append(dict(spain_series=r['series_id'],spain_name=r['name'],spain_unit=r['unit'],spain_status=r['status'],
            concept=key or '',china_series=co['selected_series'] if co else '',status=co['status'] if co else 'ADICIONAL_EXCEL_NO_EQUIVALENCIA_ACREDITADA',
            candidate_rule=json.dumps(candidate,ensure_ascii=False) if candidate else '',
            note='Correspondencia conceptual; código/parecido numérico no acredita igualdad. Fuentes, bases y universos separados.'))
    pl.DataFrame(mapping).write_csv(OUT/'correspondencias_espana_china.csv')
    pl.DataFrame([dict(concept=k,requirements=';'.join(v),reuse_count=len(v),decision='Un input reutilizado; no descargar ni contar de nuevo por figura') for k,v in refs.items()]).write_csv(OUT/'inventario_requisitos.csv')
    reconciled=pl.read_csv(ROOT/'output/revision_espana/filas_reconciliadas.csv',infer_schema_length=None)
    reconciled.write_csv(OUT/'inventario_filas_excel_anotadas.csv')

if __name__=='__main__':
    d,c=build();print({'observations':d.height,'series':c.height,'analytic_series':d.filter(pl.col('analytic'))['series_id'].n_unique()})
